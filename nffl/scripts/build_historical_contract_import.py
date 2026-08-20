#!/usr/bin/env python3

from __future__ import annotations

import csv
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]

SOURCE = (
    ROOT
    / "data"
    / "historical_contracts"
    / "source"
    / "nffl_contract_history_source.xlsx"
)

GENERATED = (
    ROOT
    / "data"
    / "historical_contracts"
    / "generated"
)

SEED_SQL = ROOT / "sql" / "021_seed_nffl_historical_contracts.sql"

LEAGUE_KEY = "470.l.84346"

TEAM_KEYS = {
    "Steady": "470.l.84346.t.1",
    "Miles": "470.l.84346.t.2",
    "Justin": "470.l.84346.t.3",
    "Cole": "470.l.84346.t.4",
    "Brent": "470.l.84346.t.5",
    "Victor": "470.l.84346.t.6",
    "Zachary": "470.l.84346.t.7",
    "Tyler": "470.l.84346.t.8",
    "Chad B": "470.l.84346.t.9",
    "Christopher": "470.l.84346.t.10",
    "George": "470.l.84346.t.11",
    "Michael": "470.l.84346.t.12",
}

MANAGER_SHEETS = list(TEAM_KEYS.keys())

# Obvious source-name variants. Original source spelling is preserved
# in source_note / audit output.
CANONICAL_NAMES = {
    "Ceedee Lamb": "CeeDee Lamb",
    "TJ Hockenson": "T.J. Hockenson",
    "Keenan  Allen": "Keenan Allen",
    "Javante Williams": "Javonte Williams",
    "Jonathon Taylor": "Jonathan Taylor",
}

# Spreadsheet legend colors.
COLOR_EXPIRED = "FF999999"
COLOR_NO_CONTRACT = "FFF4CCCC"
COLOR_FT = "FF6D9EEB"
COLOR_TRADE = "FF93C47D"
COLOR_DROPPED = "FFE06666"
COLOR_WAIVER = "FFFF9900"

SEMANTIC_COLORS = {
    COLOR_EXPIRED,
    COLOR_NO_CONTRACT,
    COLOR_FT,
    COLOR_TRADE,
    COLOR_DROPPED,
    COLOR_WAIVER,
}

# User-confirmed source corrections.
#
# Keys are:
#     (sheet, actual spreadsheet row, season)
#
# Raw workbook value is NEVER overwritten in the audit.
OVERRIDES = {
    ("Miles", 9, 2025): {
        "years": 2,
        "reason": (
            "User confirmed Nick Chubb 2025 #VALUE! should be 2."
        ),
    },
    ("Tyler", 11, 2023): {
        "years": 2,
        "reason": (
            "User confirmed James Cook 2023 value 2 was accidentally erased."
        ),
    },
    ("Tyler", 11, 2024): {
        "years": 1,
        "reason": (
            "User confirmed James Cook 2024 value is 1."
        ),
    },
}


def fill_rgb(cell) -> str:
    fg = cell.fill.fgColor

    if fg.type == "rgb" and isinstance(fg.rgb, str):
        return fg.rgb.upper()

    return ""


def sql_quote(value) -> str:
    if value is None:
        return "NULL"

    return "'" + str(value).replace("'", "''") + "'"


def meaningful_source_row(ws_formula, ws_values, row: int) -> bool:
    player = ws_values.cell(row, 1).value

    if player is None:
        player = ws_formula.cell(row, 1).value

    if player is None or not str(player).strip():
        return False

    for col in range(2, 7):
        formula_value = ws_formula.cell(row, col).value
        cached_value = ws_values.cell(row, col).value
        rgb = fill_rgb(ws_formula.cell(row, col))

        if (
            formula_value is not None
            or cached_value is not None
            or rgb in SEMANTIC_COLORS
        ):
            return True

    return False


def classify_value(
    *,
    sheet: str,
    row: int,
    season: int,
    formula_value,
    cached_value,
    rgb: str,
):
    override = OVERRIDES.get((sheet, row, season))

    if override is not None:
        return {
            "status": "CONTRACT",
            "years": override["years"],
            "acquisition": "NONE",
            "rule": "USER_OVERRIDE",
            "note": override["reason"],
        }

    acquisition = "NONE"

    if rgb == COLOR_WAIVER:
        acquisition = "WAIVER"
    elif rgb == COLOR_TRADE:
        acquisition = "TRADE"

    value = cached_value

    if isinstance(value, bool):
        pass

    elif isinstance(value, (int, float)):
        if float(value).is_integer():
            years = int(value)

            if 1 <= years <= 4:
                return {
                    "status": "CONTRACT",
                    "years": years,
                    "acquisition": acquisition,
                    "rule": "NUMERIC_CONTRACT",
                    "note": "",
                }

            return {
                "status": "UNRESOLVED",
                "years": None,
                "acquisition": acquisition,
                "rule": "NUMERIC_OUT_OF_RANGE",
                "note": f"Unexpected numeric value: {value}",
            }

    elif isinstance(value, str):
        text = value.strip()

        if text.upper() == "FT":
            return {
                "status": "FT",
                "years": None,
                "acquisition": "NONE",
                "rule": "FT_VALUE",
                "note": "",
            }

        if text.lower().startswith("dropped"):
            return {
                "status": "DROPPED",
                "years": None,
                "acquisition": "NONE",
                "rule": "DROPPED_VALUE",
                "note": "",
            }

        contract_match = re.match(
            r"^\s*([1-4])(?:\D.*)?$",
            text,
        )

        if contract_match:
            return {
                "status": "CONTRACT",
                "years": int(contract_match.group(1)),
                "acquisition": acquisition,
                "rule": "TEXT_CONTRACT",
                "note": "",
            }

        return {
            "status": "UNRESOLVED",
            "years": None,
            "acquisition": acquisition,
            "rule": "UNRECOGNIZED_TEXT",
            "note": f"Unexpected text: {text}",
        }

    # Blank cells still carry spreadsheet semantics through color.
    if rgb == COLOR_DROPPED:
        return {
            "status": "DROPPED",
            "years": None,
            "acquisition": "NONE",
            "rule": "RED_DROP",
            "note": "Drop indicated by spreadsheet fill color.",
        }

    if rgb == COLOR_EXPIRED:
        return {
            "status": "EXPIRED",
            "years": None,
            "acquisition": "NONE",
            "rule": "GRAY_EXPIRED",
            "note": "Expiration indicated by spreadsheet fill color.",
        }

    if rgb == COLOR_NO_CONTRACT:
        return {
            "status": "PINK_PENDING",
            "years": None,
            "acquisition": "NONE",
            "rule": "PINK_NO_CONTRACT",
            "note": "",
        }

    return {
        "status": "BLANK",
        "years": None,
        "acquisition": "NONE",
        "rule": "BLANK",
        "note": "",
    }


def manager_team_names(wb_values):
    ws = wb_values["Managers"]

    result = {}

    for row in range(2, ws.max_row + 1):
        owner = ws.cell(row, 2).value
        team = ws.cell(row, 3).value

        if owner and team:
            result[str(owner).strip()] = str(team).strip()

    return result


def main() -> None:
    if not SOURCE.exists():
        raise SystemExit(f"Missing workbook: {SOURCE}")

    GENERATED.mkdir(parents=True, exist_ok=True)

    wb_formula = load_workbook(
        SOURCE,
        data_only=False,
        read_only=False,
    )

    wb_values = load_workbook(
        SOURCE,
        data_only=True,
        read_only=False,
    )

    missing_sheets = [
        sheet
        for sheet in MANAGER_SHEETS
        if sheet not in wb_formula.sheetnames
    ]

    if missing_sheets:
        raise SystemExit(
            "Missing expected manager sheets: "
            + ", ".join(missing_sheets)
        )

    team_names = manager_team_names(wb_values)

    episodes = []
    seasons = []
    audit = []
    blocking = []

    for sheet in MANAGER_SHEETS:
        ws_formula = wb_formula[sheet]
        ws_values = wb_values[sheet]

        years = [
            int(ws_values.cell(2, col).value)
            for col in range(2, 7)
        ]

        if years != [2025, 2024, 2023, 2022, 2021]:
            raise SystemExit(
                f"{sheet}: unexpected season columns: {years}"
            )

        team_name = team_names.get(sheet)

        if not team_name:
            raise SystemExit(
                f"{sheet}: team name missing from Managers tab."
            )

        for row in range(3, ws_formula.max_row + 1):
            if not meaningful_source_row(
                ws_formula,
                ws_values,
                row,
            ):
                continue

            raw_player = ws_values.cell(row, 1).value

            if raw_player is None:
                raw_player = ws_formula.cell(row, 1).value

            raw_player = str(raw_player).strip()

            player = CANONICAL_NAMES.get(
                raw_player,
                raw_player,
            )

            source_row_number = row - 2

            source_note_parts = [
                f"Workbook source: {sheet}!A{row}"
            ]

            if player != raw_player:
                source_note_parts.append(
                    f"Source player name normalized from "
                    f"'{raw_player}' to '{player}'"
                )

                audit.append({
                    "category": "PLAYER_NAME_NORMALIZED",
                    "source_sheet": sheet,
                    "source_row": row,
                    "player_name_raw": raw_player,
                    "season": "",
                    "detail": f"{raw_player} -> {player}",
                    "blocking": "FALSE",
                })

            episodes.append({
                "league_key": LEAGUE_KEY,
                "team_key": TEAM_KEYS[sheet],
                "source_row_number": source_row_number,
                "source_owner_name": sheet,
                "source_team_name": team_name,
                "player_name": player,
                "yahoo_player_key": "",
                "source_note": "; ".join(source_note_parts),
            })

            interpreted = {}

            for col, season in zip(
                range(2, 7),
                years,
            ):
                formula_value = ws_formula.cell(
                    row,
                    col,
                ).value

                cached_value = ws_values.cell(
                    row,
                    col,
                ).value

                rgb = fill_rgb(
                    ws_formula.cell(row, col)
                )

                result = classify_value(
                    sheet=sheet,
                    row=row,
                    season=season,
                    formula_value=formula_value,
                    cached_value=cached_value,
                    rgb=rgb,
                )

                result["formula_value"] = formula_value
                result["cached_value"] = cached_value
                result["fill_rgb"] = rgb

                interpreted[season] = result

            # Resolve pink "No Contract" cells chronologically.
            #
            # If a player had 1 year remaining immediately before
            # the pink cell, that first pink year is the expiration.
            #
            # If >1 year remained, the contract ended early, which
            # represents a drop.
            #
            # Otherwise the pink cell remains NO_CONTRACT.
            previous = None

            for season in sorted(years):
                result = interpreted[season]

                if result["status"] == "PINK_PENDING":
                    if (
                        previous
                        and previous["status"] == "CONTRACT"
                        and previous["years"] == 1
                    ):
                        result["status"] = "EXPIRED"
                        result["rule"] = "INFERRED_NATURAL_EXPIRATION"
                        result["note"] = (
                            "Pink No Contract cell immediately "
                            "follows a 1-year contract."
                        )

                        audit.append({
                            "category": "INFERRED_EXPIRATION",
                            "source_sheet": sheet,
                            "source_row": row,
                            "player_name_raw": raw_player,
                            "season": season,
                            "detail": result["note"],
                            "blocking": "FALSE",
                        })

                    elif (
                        previous
                        and previous["status"] == "CONTRACT"
                        and previous["years"] is not None
                        and previous["years"] > 1
                    ):
                        result["status"] = "DROPPED"
                        result["rule"] = "INFERRED_EARLY_DROP"
                        result["note"] = (
                            "Pink No Contract cell follows a "
                            f"{previous['years']}-year contract; "
                            "contract ended before natural expiration."
                        )

                        audit.append({
                            "category": "INFERRED_EARLY_DROP",
                            "source_sheet": sheet,
                            "source_row": row,
                            "player_name_raw": raw_player,
                            "season": season,
                            "detail": result["note"],
                            "blocking": "FALSE",
                        })

                    else:
                        result["status"] = "NO_CONTRACT"
                        result["rule"] = "NO_CONTRACT"
                        result["note"] = (
                            "No Contract indicated by spreadsheet "
                            "fill color."
                        )

                if result["status"] != "BLANK":
                    previous = result

            for season in years:
                result = interpreted[season]

                raw_source = result["cached_value"]

                if raw_source is None:
                    raw_source_text = ""
                else:
                    raw_source_text = str(raw_source)

                audit_row = {
                    "category": result["rule"],
                    "source_sheet": sheet,
                    "source_row": row,
                    "player_name_raw": raw_player,
                    "season": season,
                    "detail": result["note"],
                    "blocking": (
                        "TRUE"
                        if result["status"] == "UNRESOLVED"
                        else "FALSE"
                    ),
                }

                if result["rule"] == "USER_OVERRIDE":
                    audit.append({
                        **audit_row,
                        "category": "USER_CONFIRMED_OVERRIDE",
                    })

                if result["status"] == "UNRESOLVED":
                    audit.append(audit_row)
                    blocking.append(audit_row)

                if result["status"] in {
                    "BLANK",
                    "UNRESOLVED",
                }:
                    continue

                seasons.append({
                    "league_key": LEAGUE_KEY,
                    "team_key": TEAM_KEYS[sheet],
                    "source_row_number": source_row_number,
                    "season_year": season,
                    "contract_years": (
                        result["years"]
                        if result["years"] is not None
                        else ""
                    ),
                    "contract_status": result["status"],
                    "acquisition_type": result["acquisition"],
                    "source_value": raw_source_text,
                    "note": result["note"],
                })

    if blocking:
        for item in blocking:
            print(
                "BLOCKING:",
                item["source_sheet"],
                item["source_row"],
                item["player_name_raw"],
                item["season"],
                item["detail"],
            )

        raise SystemExit(
            f"{len(blocking)} unresolved source cells remain."
        )

    if len(episodes) != 196:
        raise SystemExit(
            f"Expected 196 episodes; found {len(episodes)}."
        )

    episode_keys = {
        (
            row["league_key"],
            row["team_key"],
            row["source_row_number"],
        )
        for row in episodes
    }

    if len(episode_keys) != len(episodes):
        raise SystemExit(
            "Duplicate episode primary keys detected."
        )

    season_keys = {
        (
            row["league_key"],
            row["team_key"],
            row["source_row_number"],
            row["season_year"],
        )
        for row in seasons
    }

    if len(season_keys) != len(seasons):
        raise SystemExit(
            "Duplicate season primary keys detected."
        )

    allowed_statuses = {
        "CONTRACT",
        "FT",
        "NO_CONTRACT",
        "EXPIRED",
        "DROPPED",
    }

    allowed_acquisition = {
        "NONE",
        "TRADE",
        "WAIVER",
    }

    for row in seasons:
        if row["contract_status"] not in allowed_statuses:
            raise SystemExit(
                f"Invalid status: {row}"
            )

        if row["acquisition_type"] not in allowed_acquisition:
            raise SystemExit(
                f"Invalid acquisition: {row}"
            )

        if row["contract_status"] == "CONTRACT":
            years = int(row["contract_years"])

            if years not in {1, 2, 3, 4}:
                raise SystemExit(
                    f"Invalid contract years: {row}"
                )

        elif row["contract_years"] != "":
            raise SystemExit(
                f"Non-contract row contains years: {row}"
            )

    episodes.sort(
        key=lambda row: (
            int(row["team_key"].split(".t.")[-1]),
            int(row["source_row_number"]),
        )
    )

    seasons.sort(
        key=lambda row: (
            int(row["team_key"].split(".t.")[-1]),
            int(row["source_row_number"]),
            int(row["season_year"]),
        )
    )

    audit.sort(
        key=lambda row: (
            row["source_sheet"],
            int(row["source_row"]),
            int(row["season"]) if row["season"] != "" else 0,
            row["category"],
        )
    )

    episode_fields = [
        "league_key",
        "team_key",
        "source_row_number",
        "source_owner_name",
        "source_team_name",
        "player_name",
        "yahoo_player_key",
        "source_note",
    ]

    season_fields = [
        "league_key",
        "team_key",
        "source_row_number",
        "season_year",
        "contract_years",
        "contract_status",
        "acquisition_type",
        "source_value",
        "note",
    ]

    audit_fields = [
        "category",
        "source_sheet",
        "source_row",
        "player_name_raw",
        "season",
        "detail",
        "blocking",
    ]

    def write_csv(path, fields, rows):
        with path.open(
            "w",
            encoding="utf-8-sig",
            newline="",
        ) as handle:
            writer = csv.DictWriter(
                handle,
                fieldnames=fields,
            )
            writer.writeheader()
            writer.writerows(rows)

    write_csv(
        GENERATED / "episodes.csv",
        episode_fields,
        episodes,
    )

    write_csv(
        GENERATED / "seasons.csv",
        season_fields,
        seasons,
    )

    write_csv(
        GENERATED / "validation.csv",
        audit_fields,
        audit,
    )

    sql = []

    sql.append(
        "-- Generated by scripts/build_historical_contract_import.py"
    )
    sql.append(
        "-- Source: local nffl_contract_history_source.xlsx"
    )
    sql.append(
        "-- Includes user-confirmed Nick Chubb and James Cook corrections."
    )
    sql.append("")

    sql.append(
        "INSERT INTO nffl.historical_contract_episode ("
    )
    sql.append(
        "    league_key,"
    )
    sql.append(
        "    team_key,"
    )
    sql.append(
        "    source_row_number,"
    )
    sql.append(
        "    source_owner_name,"
    )
    sql.append(
        "    source_team_name,"
    )
    sql.append(
        "    player_name,"
    )
    sql.append(
        "    yahoo_player_key,"
    )
    sql.append(
        "    source_note"
    )
    sql.append(
        ") VALUES"
    )

    episode_values = []

    for row in episodes:
        episode_values.append(
            "    ("
            + ", ".join([
                sql_quote(row["league_key"]),
                sql_quote(row["team_key"]),
                str(row["source_row_number"]),
                sql_quote(row["source_owner_name"]),
                sql_quote(row["source_team_name"]),
                sql_quote(row["player_name"]),
                "NULL",
                sql_quote(row["source_note"]),
            ])
            + ")"
        )

    sql.append(",\n".join(episode_values))

    sql.append(
        "ON CONFLICT (league_key, team_key, source_row_number)"
    )
    sql.append(
        "DO UPDATE SET"
    )
    sql.append(
        "    source_owner_name = EXCLUDED.source_owner_name,"
    )
    sql.append(
        "    source_team_name = EXCLUDED.source_team_name,"
    )
    sql.append(
        "    player_name = EXCLUDED.player_name,"
    )
    sql.append(
        "    yahoo_player_key = COALESCE("
    )
    sql.append(
        "        EXCLUDED.yahoo_player_key,"
    )
    sql.append(
        "        nffl.historical_contract_episode.yahoo_player_key"
    )
    sql.append(
        "    ),"
    )
    sql.append(
        "    source_note = EXCLUDED.source_note;"
    )
    sql.append("")

    sql.append(
        "INSERT INTO nffl.historical_contract_season ("
    )
    sql.append(
        "    league_key,"
    )
    sql.append(
        "    team_key,"
    )
    sql.append(
        "    source_row_number,"
    )
    sql.append(
        "    season_year,"
    )
    sql.append(
        "    contract_years,"
    )
    sql.append(
        "    contract_status,"
    )
    sql.append(
        "    acquisition_type,"
    )
    sql.append(
        "    source_value,"
    )
    sql.append(
        "    note"
    )
    sql.append(
        ") VALUES"
    )

    season_values = []

    for row in seasons:
        contract_years = (
            str(row["contract_years"])
            if row["contract_years"] != ""
            else "NULL"
        )

        season_values.append(
            "    ("
            + ", ".join([
                sql_quote(row["league_key"]),
                sql_quote(row["team_key"]),
                str(row["source_row_number"]),
                str(row["season_year"]),
                contract_years,
                sql_quote(row["contract_status"]),
                sql_quote(row["acquisition_type"]),
                sql_quote(row["source_value"]),
                sql_quote(row["note"]),
            ])
            + ")"
        )

    sql.append(",\n".join(season_values))

    sql.append(
        "ON CONFLICT ("
        "league_key, team_key, source_row_number, season_year"
        ")"
    )
    sql.append(
        "DO UPDATE SET"
    )
    sql.append(
        "    contract_years = EXCLUDED.contract_years,"
    )
    sql.append(
        "    contract_status = EXCLUDED.contract_status,"
    )
    sql.append(
        "    acquisition_type = EXCLUDED.acquisition_type,"
    )
    sql.append(
        "    source_value = EXCLUDED.source_value,"
    )
    sql.append(
        "    note = EXCLUDED.note;"
    )
    sql.append("")

    SEED_SQL.write_text(
        "\n".join(sql),
        encoding="utf-8",
        newline="\n",
    )

    statuses = Counter(
        row["contract_status"]
        for row in seasons
    )

    acquisitions = Counter(
        row["acquisition_type"]
        for row in seasons
    )

    teams = Counter(
        row["team_key"]
        for row in episodes
    )

    print(f"EPISODES={len(episodes)}")
    print(f"SEASONS={len(seasons)}")
    print(f"TEAMS={len(teams)}")

    for status in sorted(statuses):
        print(
            f"STATUS_{status}={statuses[status]}"
        )

    for acquisition in sorted(acquisitions):
        print(
            f"ACQUISITION_{acquisition}="
            f"{acquisitions[acquisition]}"
        )

    print(
        "INFERRED_EARLY_DROPS="
        + str(
            sum(
                1
                for row in audit
                if row["category"]
                == "INFERRED_EARLY_DROP"
            )
        )
    )

    print(
        "USER_OVERRIDES="
        + str(
            sum(
                1
                for row in audit
                if row["category"]
                == "USER_CONFIRMED_OVERRIDE"
            )
        )
    )

    print("BLOCKING_ISSUES=0")
    print(f"SEED_SQL={SEED_SQL}")
    print("VALIDATION=PASS")


if __name__ == "__main__":
    main()